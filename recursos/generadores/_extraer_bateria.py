# -*- coding: utf-8 -*-
"""Extrae los nueve instrumentos del libro de evaluacion a contenido/bateria.py.

Cada hoja tiene su propia disposicion, asi que hay un esquema por instrumento.
Se ejecuta una vez; despues el generador del formulario lee el modulo, no el
Excel. Los enunciados no se retoclean a mano en ningun sitio.
"""
import os, re, sys, json, warnings
warnings.filterwarnings('ignore')
import openpyxl

AQUI = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
LIBRO = '/Users/javier/Downloads/Evaluación Psicológica — Hogar Terapéutico.xlsx'
DESTINO = os.path.join(AQUI, 'contenido', 'bateria.py')
col = openpyxl.utils.column_index_from_string
limpio = lambda v: re.sub(r'[ \t]+', ' ', str(v).strip()) if v is not None else ''

ESCALAS = {
    'likert5': ['Muy en desacuerdo', 'Ligeramente en desacuerdo',
                'Ni de acuerdo, ni en desacuerdo', 'Ligeramente de acuerdo',
                'Muy de acuerdo'],
    'stai':    ['Nada', 'Algo', 'Bastante', 'Mucho'],
    'frec5':   ['Casi nunca', 'Pocas veces', 'Unas veces sí y otras no',
                'Muchas veces', 'Casi siempre'],
}


def numerados(ws, cn, ct, f0, paso=1):
    out = []
    for f in range(f0, ws.max_row + 1, paso):
        n, t = ws.cell(f, cn).value, limpio(ws.cell(f, ct).value)
        if n is None or not t:
            continue
        try:
            n = int(float(n))
        except (TypeError, ValueError):
            continue
        out.append(t)
    return out


def columnas(ws, letras, f0, paso=2):
    """El IA-TP reparte sus adjetivos en varias columnas, de dos en dos filas."""
    out = []
    for L in letras:
        for f in range(f0, ws.max_row + 1, paso):
            v = limpio(ws.cell(f, col(L)).value)
            if v and len(v) < 26 and ' ' not in v:
                out.append(v)
    return out


def bdi(ws, c, f0, f1):
    """Cada celda trae el grupo entero: titulo y opciones separadas por saltos."""
    grupos = []
    for f in range(f0, f1 + 1):
        v = ws.cell(f, c).value
        if not v:
            continue
        lineas = [re.sub(r'\s+', ' ', x).strip() for x in str(v).split('\n') if x.strip()]
        cabecera = re.sub(r'^\s*\d{1,2}\s*\.\s*', '', lineas[0]).strip()
        opciones = []
        for x in lineas[1:]:
            m = re.match(r'^(\d[ab]?)\s*\.?\s+(.*)$', x)
            if m:
                opciones.append(f'{m.group(1)} · {m.group(2)}')
            elif opciones:
                opciones[-1] += ' ' + x
        if cabecera and opciones:
            grupos.append({'titulo': cabecera, 'opciones': opciones})
    return grupos


def creencias(ws, c, f0):
    """Los encabezados de categoria van en negrita y con relleno. Adivinarlos por
    como estan escritos fallaba: "Merezco morir" parece un titulo y es un item."""
    out, cat = [], None
    for f in range(f0, ws.max_row + 1):
        celda = ws.cell(f, c)
        v = limpio(celda.value)
        if not v:
            continue
        if celda.font and celda.font.bold:
            cat = v
            continue
        out.append({'categoria': cat, 'texto': v})
    return out


wb = openpyxl.load_workbook(LIBRO)
B = {}

B['5big'] = dict(
    rotulo='Personalidad (5BIG)', tipo='escala', escala='likert5',
    enunciado='Las siguientes expresiones le describen a usted con más o menos precisión. '
              'Marque la opción que considere correcta.',
    items=numerados(wb['1. Personalidad'], col('D'), col('E'), 8))

B['af5'] = dict(
    rotulo='Autoconcepto (AF-5)', tipo='numero', minimo=1, maximo=99,
    enunciado='Conteste con un valor entre 1 y 99 según su grado de acuerdo con cada frase. '
              '1 es total desacuerdo y 99 total acuerdo.',
    items=numerados(wb['2. Autoconcepto'], col('D'), col('E'), 8))

B['stai_e'] = dict(
    rotulo='Ansiedad: cómo se siente AHORA (STAI-E)', tipo='escala', escala='stai',
    enunciado='Señale lo que indique mejor cómo se siente usted AHORA MISMO, en este momento.',
    items=numerados(wb['3. Ansiedad Estado'], col('C'), col('D'), 8))

B['stai_r'] = dict(
    rotulo='Ansiedad: cómo se siente EN GENERAL (STAI-R)', tipo='escala', escala='stai',
    enunciado='Señale lo que indique mejor cómo se siente usted EN GENERAL, en la mayoría '
              'de las ocasiones.',
    items=numerados(wb['3. Ansiedad Rasgo'], col('C'), col('D'), 9))

B['hipno'] = dict(
    rotulo='Sugestionabilidad', tipo='escala', escala='frec5',
    enunciado='Valore la frecuencia con que se producen estas situaciones en su vida cotidiana.',
    items=numerados(wb['4. Hipno'], col('D'), col('E'), 8))

B['apego'] = dict(
    rotulo='Modelos de apego', tipo='lista',
    enunciado='Marque solo aquellas afirmaciones que considere correctas en su caso.',
    items=numerados(wb['5. Apego'], col('D'), col('E'), 8))

B['aron'] = dict(
    rotulo='Alta sensibilidad (Aron)', tipo='lista',
    enunciado='Marque las afirmaciones con las que se sienta identificado/a. '
              'Si tiene dudas, no la marque.',
    items=numerados(wb['6. Sensibilidad'], col('D'), col('E'), 9))

B['bdi'] = dict(
    rotulo='Estado de ánimo (BDI-II)', tipo='grupos',
    enunciado='En cada grupo, elija la frase que mejor describa cómo se ha sentido las últimas '
              'dos semanas, incluyendo hoy. Si varias le parecen igual de apropiadas, marque '
              'la del número más alto.',
    items=bdi(wb['7. Depresion'], col('D'), 8, 28))

B['creencias'] = dict(
    rotulo='Creencias negativas', tipo='lista_cat',
    enunciado='Marque aquellas creencias negativas con las que cree que se identifica.',
    items=creencias(wb['8. Creencias'], col('C'), 8))

B['iatp'] = dict(
    rotulo='Adjetivos que le definen (IA-TP)', tipo='lista',
    enunciado='Marque los adjetivos que cree que le definen mejor.',
    items=columnas(wb['9. Rasgos IATP'], ['G', 'O', 'X', 'AG', 'AP'], 8))

total = sum(len(v['items']) for v in B.values())
for k, v in B.items():
    print(f"  {v['rotulo'][:42]:44} {v['tipo']:10} {len(v['items']):>4}")
print(f"  {'TOTAL':44} {'':10} {total:>4}")

cab = '''# -*- coding: utf-8 -*-
"""Los nueve instrumentos del libro de evaluacion, extraidos de sus hojas.

GENERADO por generadores/_extraer_bateria.py. No editar a mano.
Fuente: Evaluacion Psicologica - Hogar Terapeutico.xlsx
"""

ESCALAS = '''
# json.dumps escribe null/true/false, que no son literales de Python
import pprint
with open(DESTINO, 'w', encoding='utf-8') as f:
    f.write(cab + pprint.pformat(ESCALAS, width=96, sort_dicts=False) + '\n\nBATERIA = '
            + pprint.pformat(B, width=96, sort_dicts=False) + '\n')
print(f'\n  escrito {os.path.relpath(DESTINO, AQUI)}')
